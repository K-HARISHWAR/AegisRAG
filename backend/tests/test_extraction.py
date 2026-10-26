"""Unit tests for PDF, DOCX, and XLSX safe text extraction and provenance tracking."""

import io
import pytest
import pypdf
import docx
import openpyxl

from app.services.extraction import (
    extract_document_text,
    extract_docx,
    extract_pdf,
    extract_xlsx,
)


def _create_sample_pdf(pages: list[str]) -> bytes:
    """Helper to generate a valid PDF binary in memory with specified page texts."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text in pages:
        c.drawString(72, 720, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def _create_sample_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """Helper to generate a valid DOCX binary in memory."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)

    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, val in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = val

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _create_sample_xlsx(sheets: dict[str, list[list[str]]]) -> bytes:
    """Helper to generate a valid XLSX binary in memory."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_pdf_extraction_page_provenance() -> None:
    """Verify PDF extraction preserves page numbers and text content accurately."""
    page_texts = [
        "First page content: System architecture overview and deployment protocols.",
        "Second page content: Production operational runbook and emergency escalation contacts.",
    ]
    pdf_bytes = _create_sample_pdf(page_texts)

    result = extract_pdf(pdf_bytes)
    assert result.status == "success"
    assert len(result.passages) == 2
    assert result.passages[0].location == "Page 1"
    assert "System architecture" in result.passages[0].text
    assert result.passages[1].location == "Page 2"
    assert "escalation contacts" in result.passages[1].text
    assert result.total_characters > 50


def test_scanned_or_empty_pdf_warning() -> None:
    """Verify empty or scanned PDFs trigger scanned_warning status rather than false success."""
    # Blank PDF with no text streams
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    blank_pdf = buf.getvalue()

    result = extract_pdf(blank_pdf)
    assert result.status == "scanned_warning"
    assert "Scanned PDF detected" in result.warning_message
    assert len(result.passages) == 0


def test_corrupt_pdf_fails_safely() -> None:
    """Verify corrupt PDF binary returns status='failed' safely without raising unhandled errors."""
    corrupt_bytes = b"%PDF-corrupted-stream-with-broken-structure"
    result = extract_pdf(corrupt_bytes)
    assert result.status == "failed"
    assert result.warning_message is not None
    assert len(result.passages) == 0


def test_docx_extraction_paragraphs_and_tables() -> None:
    """Verify DOCX extraction preserves paragraph sequence and table row cells."""
    paragraphs = [
        "Engineering Infrastructure Secret Management Policy.",
        "All secrets must be stored in encrypted vaults and rotated every 90 days.",
    ]
    table_rows = [
        ["Service", "Role", "Rotation Period"],
        ["PostgreSQL", "Database Admin", "30 Days"],
        ["ChromaDB", "Vector Service", "90 Days"],
    ]
    docx_bytes = _create_sample_docx(paragraphs, table_rows)

    result = extract_docx(docx_bytes)
    assert result.status == "success"
    assert len(result.passages) >= 3

    # Paragraph verification
    assert result.passages[0].location == "Paragraph 1"
    assert "Infrastructure Secret" in result.passages[0].text
    assert result.passages[1].location == "Paragraph 2"

    # Table verification
    table_passages = [p for p in result.passages if "Table" in p.location]
    assert len(table_passages) == 3
    assert "PostgreSQL | Database Admin | 30 Days" in table_passages[1].text


def test_xlsx_extraction_sheets_and_cell_coordinates() -> None:
    """Verify XLSX extraction extracts cell data preserving sheet names and row coordinates."""
    sheet_data = {
        "ExecutiveBonuses": [
            ["Employee ID", "Department", "Bonus Target", "Award Percentage"],
            ["ENG-001", "Engineering", "25000", "110%"],
            ["HR-002", "Human Resources", "18000", "100%"],
        ],
        "Summary": [
            ["Total Pool", "Disbursed", "Remaining"],
            ["500000", "43000", "457000"],
        ]
    }
    xlsx_bytes = _create_sample_xlsx(sheet_data)

    result = extract_xlsx(xlsx_bytes)
    assert result.status == "success"
    assert len(result.passages) == 5

    # Check locations
    locations = [p.location for p in result.passages]
    assert any("Sheet: ExecutiveBonuses!Row 1" in loc for loc in locations)
    assert any("Sheet: Summary!Row 2" in loc for loc in locations)

    # Check content
    all_text = " ".join([p.text for p in result.passages])
    assert "ENG-001" in all_text
    assert "500000" in all_text


def test_extract_document_text_dispatcher() -> None:
    """Verify extract_document_text dispatches to correct parser and rejects unsupported types."""
    pdf_bytes = _create_sample_pdf([
        "Production operational runbook: comprehensive deployment and infrastructure management procedures."
    ])
    res_pdf = extract_document_text("pdf", pdf_bytes)
    assert res_pdf.status == "success"

    res_unsupported = extract_document_text("exe", b"not-a-supported-file")
    assert res_unsupported.status == "failed"
    assert "Unsupported format" in res_unsupported.warning_message
